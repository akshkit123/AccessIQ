import csv
import io
from datetime import datetime
from django.http import HttpResponse
from django.utils import timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from access_control.models import AccessLog, AccessLocation
from visitors.models import Visitor, TemporaryPass
from security.models import SecurityAlert, IssueReport


def export_access_logs_csv(queryset=None):
    if queryset is None:
        queryset = AccessLog.objects.all().select_related('location', 'student', 'staff').order_by('-timestamp')

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="access_logs_{datetime.now().strftime("%Y%m%d_%H%M")}.csv"'

    writer = csv.writer(response)
    writer.writerow([
        'Timestamp', 'Barcode/QR Scanned', 'Identity Name', 'Identity Type',
        'Location Code', 'Location Name', 'Direction', 'Barcode Verified',
        'Face Verified', 'Permission Granted', 'Final Decision', 'Failure Reason'
    ])

    for log in queryset:
        name = log.student.full_name if log.student else (log.staff.full_name if log.staff else log.visitor_name or "Unknown")
        user_type = 'Student' if log.student else ('Staff' if log.staff else ('Visitor' if log.visitor_name else 'Unknown'))
        writer.writerow([
            log.timestamp.strftime('%Y-%m-%d %H:%M:%S'),
            log.scanned_barcode,
            name,
            user_type,
            log.location.code,
            log.location.name,
            log.direction,
            'YES' if log.barcode_verified else 'NO',
            'YES' if log.face_verified else 'NO',
            'YES' if log.permission_granted else 'NO',
            log.final_decision,
            log.failure_reason
        ])

    return response


def export_access_logs_excel(queryset=None):
    if queryset is None:
        queryset = AccessLog.objects.all().select_related('location', 'student', 'staff').order_by('-timestamp')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Access Logs"

    # Styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    headers = [
        'Timestamp', 'Barcode / QR', 'User Name', 'Type',
        'Location', 'Direction', 'Face Match', 'Decision', 'Notes / Failure Reason'
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, log in enumerate(queryset, start=2):
        name = log.student.full_name if log.student else (log.staff.full_name if log.staff else log.visitor_name or "Unknown")
        user_type = 'Student' if log.student else ('Staff' if log.staff else ('Visitor' if log.visitor_name else 'Unknown'))

        row_data = [
            log.timestamp.strftime('%Y-%m-%d %H:%M:%S'),
            log.scanned_barcode,
            name,
            user_type,
            f"{log.location.code} ({log.location.name})",
            log.direction,
            'MATCH' if log.face_verified else 'MISMATCH/NONE',
            log.final_decision,
            log.failure_reason or 'Access Authorized'
        ]
        ws.append(row_data)

        # Style data row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx == 8:  # Decision
                if log.final_decision == 'GRANTED':
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="16A34A")
                else:
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="DC2626")

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="accessiq_logs_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx"'
    wb.save(response)
    return response


def export_access_logs_pdf(queryset=None):
    if queryset is None:
        queryset = AccessLog.objects.all().select_related('location', 'student', 'staff').order_by('-timestamp')[:50]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    elements = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#6B7280'),
        spaceAfter=15,
    )

    elements.append(Paragraph("AccessIQ: Smart Campus Identity, Access & Security Management", title_style))
    elements.append(Paragraph(f"Campus Access Event Audit Report • Generated: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}", subtitle_style))

    data = [[
        'Timestamp', 'Barcode/QR', 'Identity', 'Location', 'Direction', 'Face Result', 'Final Decision'
    ]]

    for log in queryset:
        name = log.student.full_name if log.student else (log.staff.full_name if log.staff else log.visitor_name or "Unknown")
        data.append([
            log.timestamp.strftime('%m/%d %H:%M'),
            log.scanned_barcode[:15],
            name[:18],
            log.location.code,
            log.direction,
            'Verified' if log.face_verified else 'No Match',
            log.final_decision
        ])

    table = Table(data, colWidths=[90, 95, 120, 85, 60, 85, 95])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (4, 0), (6, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')])
    ]))

    elements.append(table)
    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="accessiq_audit_{datetime.now().strftime("%Y%m%d_%H%M")}.pdf"'
    return response


def export_security_alerts_csv():
    alerts = SecurityAlert.objects.all().select_related('location', 'resolved_by').order_by('-timestamp')
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="security_alerts_{datetime.now().strftime("%Y%m%d_%H%M")}.csv"'

    writer = csv.writer(response)
    writer.writerow(['Timestamp', 'Alert Type', 'Severity', 'Location', 'User Identifier', 'Description', 'Status'])

    for a in alerts:
        writer.writerow([
            a.timestamp.strftime('%Y-%m-%d %H:%M:%S'),
            a.get_alert_type_display(),
            a.severity,
            a.location.name if a.location else 'Campus General',
            a.user_identifier,
            a.description,
            a.status
        ])
    return response


def export_visitors_csv():
    visitors = Visitor.objects.all().order_by('-entry_time')
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="campus_visitors_{datetime.now().strftime("%Y%m%d_%H%M")}.csv"'

    writer = csv.writer(response)
    writer.writerow(['Visitor Name', 'Host / Person Visited', 'Purpose of Visit', 'Entry Time', 'Exit Time', 'Status', 'Contact Phone'])

    for v in visitors:
        writer.writerow([
            v.full_name,
            v.person_to_meet,
            v.purpose_of_visit,
            v.entry_time.strftime('%Y-%m-%d %H:%M:%S') if v.entry_time else 'N/A',
            v.exit_time.strftime('%Y-%m-%d %H:%M:%S') if v.exit_time else 'N/A',
            v.get_status_display(),
            v.phone or ''
        ])
    return response


def export_issue_reports_csv():
    issues = IssueReport.objects.all().select_related('reporter', 'location').order_by('-created_at')
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="issue_reports_{datetime.now().strftime("%Y%m%d_%H%M")}.csv"'

    writer = csv.writer(response)
    writer.writerow(['Date / Time', 'Issue Type', 'Location', 'Reported By', 'Description', 'Status'])

    for item in issues:
        writer.writerow([
            item.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            item.get_issue_type_display(),
            item.location.name if item.location else 'Campus General',
            item.reporter.get_full_name() or item.reporter.username,
            item.description,
            item.get_status_display()
        ])
    return response
